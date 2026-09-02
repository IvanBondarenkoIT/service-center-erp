from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, func, select

from app.database import SessionLocal
from app.models import CashEntry, CashEntryKind, CashOpening, CollectionSubtype, ServiceCenter
from app.services.import_cash_xlsx import import_cash_xlsx, parse_cash_date

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "cash_book_sample.xlsx"


def _wipe_sample_period() -> None:
    db = SessionLocal()
    try:
        center_id = db.scalar(select(ServiceCenter.id).where(ServiceCenter.code == "batumi"))
        db.execute(
            delete(CashEntry).where(
                CashEntry.service_center_id == center_id,
                CashEntry.entry_date >= date(2019, 3, 1),
                CashEntry.entry_date <= date(2019, 3, 31),
            )
        )
        db.execute(
            delete(CashOpening).where(
                CashOpening.service_center_id == center_id,
                CashOpening.as_of_date == date(2019, 3, 1),
            )
        )
        db.commit()
    finally:
        db.close()


@pytest.fixture(autouse=True)
def _clean_sample_period():
    _wipe_sample_period()
    yield
    _wipe_sample_period()


def _batumi_period_entries():
    db = SessionLocal()
    try:
        center_id = db.scalar(select(ServiceCenter.id).where(ServiceCenter.code == "batumi"))
        entries = list(
            db.scalars(
                select(CashEntry).where(
                    CashEntry.service_center_id == center_id,
                    CashEntry.entry_date >= date(2019, 3, 1),
                    CashEntry.entry_date <= date(2019, 3, 31),
                )
            )
        )
        opening = db.scalar(
            select(CashOpening).where(
                CashOpening.service_center_id == center_id,
                CashOpening.as_of_date == date(2019, 3, 1),
            )
        )
        return entries, opening
    finally:
        db.close()


def test_parse_cash_date() -> None:
    assert parse_cash_date(datetime(2025, 12, 23)) == date(2025, 12, 23)
    assert parse_cash_date("25,12,25") == date(2025, 12, 25)
    assert parse_cash_date("6,01,26") == date(2026, 1, 6)
    assert parse_cash_date("07,08,26") == date(2026, 8, 7)
    assert parse_cash_date(None) is None


def test_fixture_workbook_exists_and_has_income_columns() -> None:
    assert FIXTURE.exists(), f"missing fixture {FIXTURE}"
    wb = load_workbook(FIXTURE, data_only=True)
    ws = wb.active
    assert ws["B1"].value == "Приход"
    assert ws["B3"].value == 80
    assert ws["D3"].value == 55
    assert ws["R1"].value == 200


def test_import_cash_xlsx_into_batumi() -> None:
    db = SessionLocal()
    try:
        stats = import_cash_xlsx(db, FIXTURE, center_code="batumi")
    finally:
        db.close()

    assert stats["expenses_created"] == 2
    assert stats["collections_created"] == 1
    assert stats["openings_upserted"] == 1
    assert stats["income_rows_ignored"] >= 1

    entries, opening = _batumi_period_entries()
    expenses = [e for e in entries if e.kind == CashEntryKind.expense]
    collections = [e for e in entries if e.kind == CashEntryKind.collection]
    assert len(expenses) == 2
    assert len(collections) == 1
    assert sum((e.amount for e in expenses), Decimal("0")) == Decimal("60.00")
    assert collections[0].amount == Decimal("100.00")
    assert collections[0].subtype == CollectionSubtype.bank
    assert {e.description for e in expenses} == {"fixture-heater", "fixture-keys"}
    assert opening is not None
    assert opening.amount == Decimal("200.00")
    assert opening.as_of_date == date(2019, 3, 1)


def test_import_cash_xlsx_is_idempotent() -> None:
    db = SessionLocal()
    try:
        first = import_cash_xlsx(db, FIXTURE, center_code="batumi")
        second = import_cash_xlsx(db, FIXTURE, center_code="batumi")
    finally:
        db.close()

    entries, opening = _batumi_period_entries()
    assert len(entries) == 3
    assert second["expenses_created"] == 0
    assert second["collections_created"] == 0
    assert second["expenses_skipped"] >= 2
    assert second["collections_skipped"] >= 1
    assert second["openings_upserted"] == 1
    assert opening is not None
    assert opening.amount == Decimal("200.00")
    assert first["expenses_created"] == 2
    assert first["collections_created"] == 1


def test_income_columns_do_not_create_cash_entries() -> None:
    db = SessionLocal()
    try:
        import_cash_xlsx(db, FIXTURE, center_code="batumi")
        center_id = db.scalar(select(ServiceCenter.id).where(ServiceCenter.code == "batumi"))
        by_amount = list(
            db.scalars(
                select(CashEntry).where(
                    CashEntry.service_center_id == center_id,
                    CashEntry.amount.in_((Decimal("80"), Decimal("55"), Decimal("99"), Decimal("12"))),
                    CashEntry.entry_date >= date(2019, 3, 1),
                    CashEntry.entry_date <= date(2019, 3, 31),
                )
            )
        )
        kinds = db.scalar(
            select(func.count()).select_from(CashEntry).where(
                CashEntry.service_center_id == center_id,
                CashEntry.entry_date >= date(2019, 3, 1),
                CashEntry.entry_date <= date(2019, 3, 31),
            )
        )
    finally:
        db.close()

    assert by_amount == []
    assert kinds == 3


def test_same_workbook_duplicate_rows_imported_once(tmp_path) -> None:
    from datetime import datetime

    from openpyxl import Workbook

    path = tmp_path / "dup_cash.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "March"
    ws["A1"] = "Число/ Date"
    ws["B1"] = "Приход"
    ws["H1"] = "Число"
    ws["I1"] = "Расход"
    ws["L1"] = "Число"
    ws["M1"] = "Инкасация"
    ws["Q1"] = "Касса"
    ws["R1"] = 10
    ws["B2"] = "Касса /Cash"
    ws["D2"] = "Терминал сервис"
    ws["J2"] = "На что"
    ws["N2"] = "Кому"
    ws["A3"] = datetime(2019, 3, 1)
    ws["H3"] = datetime(2019, 3, 5)
    ws["I3"] = 11
    ws["J3"] = "same-row"
    ws["H4"] = datetime(2019, 3, 5)
    ws["I4"] = 11
    ws["J4"] = "same-row"
    wb.save(path)

    db = SessionLocal()
    try:
        stats = import_cash_xlsx(db, path, center_code="batumi")
    finally:
        db.close()

    assert stats["expenses_created"] == 1
    assert stats["expenses_skipped"] == 1
    entries, _ = _batumi_period_entries()
    assert len([e for e in entries if e.description == "same-row"]) == 1

