from __future__ import annotations

import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import (
    Client,
    ClientMachine,
    LineType,
    Machine,
    OrderLine,
    OrderStatus,
    PaymentType,
    ServiceCenter,
    ServiceOrder,
    User,
    UserRole,
)
from app.services.phones import normalize_phone


def _to_decimal(value: Any) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    try:
        return Decimal(str(value).replace(",", "."))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _to_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _map_payment(raw: Any) -> PaymentType:
    text = str(raw or "").strip().lower()
    mapping = {
        "cash": PaymentType.cash,
        "card": PaymentType.card,
        "garanty": PaymentType.warranty,
        "garant": PaymentType.warranty,
        "warranty": PaymentType.warranty,
        "rasxod": PaymentType.expense,
        "perechisleniye": PaymentType.transfer,
        "perechislenie": PaymentType.transfer,
    }
    return mapping.get(text, PaymentType.cash)


def _map_line_type(payment: PaymentType, comment: str, part_name: str) -> LineType:
    c = (comment or "").lower()
    if payment == PaymentType.expense:
        return LineType.expense
    if payment == PaymentType.warranty:
        return LineType.warranty
    if "diagnost" in c or "диагност" in c:
        return LineType.diagnostic
    if part_name or "nawili" in c or "part" in c:
        return LineType.part
    return LineType.work


def _matching_order(
    db: Session,
    *,
    center_id: int,
    order_date: date,
    model_name: str,
    comment: str,
    phone: str,
    payment: PaymentType,
    cost_work: Decimal,
    cost_parts: Decimal,
) -> ServiceOrder | None:
    stmt = (
        select(ServiceOrder)
        .join(ServiceOrder.machine)
        .join(ServiceOrder.lines)
        .where(
            ServiceOrder.service_center_id == center_id,
            ServiceOrder.order_date == order_date,
            ServiceOrder.comment == comment,
            Machine.model_name == model_name,
            OrderLine.payment_type == payment,
            OrderLine.amount_work == cost_work,
            OrderLine.amount_parts == cost_parts,
        )
    )
    if phone:
        stmt = stmt.join(ServiceOrder.client).where(Client.phone == phone)
    else:
        stmt = stmt.where(ServiceOrder.client_id.is_(None))
    return db.scalar(stmt)


def import_excel(
    db: Session,
    path: Path,
    *,
    default_center_code: str = "tbilisi",
    assignee_login: str | None = None,
) -> dict[str, int]:
    settings = get_settings()
    wb = load_workbook(path, data_only=True)
    center = db.scalar(select(ServiceCenter).where(ServiceCenter.code == default_center_code))
    if not center:
        raise RuntimeError(f"Service center {default_center_code} not found; run seed first")

    if assignee_login:
        assignee = db.scalar(select(User).where(User.login == assignee_login))
    else:
        assignee = db.scalar(
            select(User).where(
                User.role == UserRole.mechanic,
                User.service_center_id == center.id,
            )
        )
    if not assignee:
        assignee = db.scalar(select(User).where(User.role == UserRole.admin))
    if not assignee:
        raise RuntimeError("No assignee user found")

    stats = {"rows": 0, "orders": 0, "skipped": 0, "clients": 0, "machines": 0, "duplicates": 0}
    seen_orders: set[tuple] = set()

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c).strip() if c is not None else "" for c in rows[0]]
        # Expected columns by index from Excel
        for raw in rows[1:]:
            stats["rows"] += 1
            order_date = _to_date(raw[0] if len(raw) > 0 else None)
            brand_model = str(raw[2] if len(raw) > 2 and raw[2] is not None else "").strip()
            comment = str(raw[3] if len(raw) > 3 and raw[3] is not None else "").strip()
            owner = str(raw[4] if len(raw) > 4 and raw[4] is not None else "").strip()
            payment_raw = raw[5] if len(raw) > 5 else None
            phone_raw = raw[6] if len(raw) > 6 else None
            # Handle August sheet quirk where Owner holds phone
            if owner and owner.replace(".", "").isdigit() and not phone_raw:
                phone_raw = owner
                owner = ""
            total = _to_decimal(raw[7] if len(raw) > 7 else None)
            cost_work = _to_decimal(raw[8] if len(raw) > 8 else None)
            cost_parts = _to_decimal(raw[9] if len(raw) > 9 else None)
            part_name = str(raw[10] if len(raw) > 10 and raw[10] is not None else "").strip()
            part_serial = str(raw[11] if len(raw) > 11 and raw[11] is not None else "").strip()

            if not order_date and not brand_model and not comment and not phone_raw:
                stats["skipped"] += 1
                continue
            if not order_date:
                stats["skipped"] += 1
                continue

            phone = normalize_phone(str(phone_raw) if phone_raw is not None else "")
            client = None
            if phone:
                client = db.scalar(select(Client).where(Client.phone == phone))
                if not client:
                    client = Client(phone=phone, name=owner or phone)
                    db.add(client)
                    db.flush()
                    stats["clients"] += 1
                elif owner and not client.name:
                    client.name = owner

            model_name = brand_model or "unknown"
            payment = _map_payment(payment_raw)
            line_type = _map_line_type(payment, comment, part_name)
            if cost_work == 0 and cost_parts == 0 and total:
                cost_work = total

            fingerprint = (
                center.id,
                order_date,
                model_name,
                comment,
                phone,
                payment.value,
                cost_work,
                cost_parts,
            )
            if fingerprint in seen_orders or _matching_order(
                db,
                center_id=center.id,
                order_date=order_date,
                model_name=model_name,
                comment=comment,
                phone=phone,
                payment=payment,
                cost_work=cost_work,
                cost_parts=cost_parts,
            ):
                stats["duplicates"] += 1
                stats["skipped"] += 1
                continue
            seen_orders.add(fingerprint)

            placeholder = (
                f"{settings.import_placeholder_serial_prefix}"
                f"{sheet_name}-{order_date.isoformat()}-{uuid.uuid4().hex[:10]}"
            )
            machine = Machine(
                serial_number=placeholder,
                model_name=model_name,
                needs_serial=True,
            )
            db.add(machine)
            db.flush()
            stats["machines"] += 1

            if client:
                link = db.scalar(
                    select(ClientMachine).where(
                        ClientMachine.client_id == client.id,
                        ClientMachine.machine_id == machine.id,
                    )
                )
                if not link:
                    db.add(ClientMachine(client_id=client.id, machine_id=machine.id))

            order = ServiceOrder(
                order_date=order_date,
                assignee_id=assignee.id,
                service_center_id=center.id,
                machine_id=machine.id,
                client_id=client.id if client else None,
                comment=comment,
                status=OrderStatus.issued,
            )
            order.lines.append(
                OrderLine(
                    line_type=line_type,
                    description=comment or part_name or model_name,
                    part_code=part_serial or part_name,
                    payment_type=payment,
                    amount_work=cost_work,
                    amount_parts=cost_parts,
                )
            )
            db.add(order)
            stats["orders"] += 1

    db.commit()
    return stats
