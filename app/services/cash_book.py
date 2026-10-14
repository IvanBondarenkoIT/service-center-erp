from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload

from app.models import (
    CashEntry,
    CashEntryKind,
    CashOpening,
    OrderLine,
    OrderStatus,
    PaymentType,
    ServiceCenter,
    ServiceOrder,
)

BUSINESS_TZ = timezone(timedelta(hours=4))  # Georgia (no DST)


def business_today() -> date:
    return datetime.now(BUSINESS_TZ).date()


def as_business_date(value: datetime | date | None) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        dt = value
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(BUSINESS_TZ).date()
    if isinstance(value, date):
        return value
    return None


def cash_entry_key(
    center_id: int,
    entry_date: date,
    kind: CashEntryKind,
    amount: Decimal | int | str | None,
    description: str | None,
) -> tuple:
    return (
        center_id,
        entry_date,
        kind.value if isinstance(kind, CashEntryKind) else str(kind),
        money(amount),
        (description or "").strip(),
    )


def existing_cash_entry(
    db: Session,
    *,
    center_id: int,
    entry_date: date,
    kind: CashEntryKind,
    amount: Decimal | int | str | None,
    description: str | None,
) -> CashEntry | None:
    key = cash_entry_key(center_id, entry_date, kind, amount, description)
    return db.scalar(
        select(CashEntry).where(
            CashEntry.service_center_id == key[0],
            CashEntry.entry_date == key[1],
            CashEntry.kind == kind,
            CashEntry.amount == key[3],
            CashEntry.description == key[4],
        )
    )

ZERO = Decimal("0.00")
CENTS = Decimal("0.01")


def money(value: Decimal | int | str | None) -> Decimal:
    return Decimal(str(value or 0)).quantize(CENTS)


# Turnover (вал). A new client payment method is one more entry here.
VAL_METHODS = (
    PaymentType.cash,
    PaymentType.card,
    PaymentType.transfer,
    PaymentType.promo,
)


@dataclass
class IncomeRow:
    order_id: int
    order_date: date
    payment_type: PaymentType
    amount: Decimal
    assignee: str
    note: str


@dataclass
class OutflowRow:
    entry_id: int
    entry_date: date
    amount: Decimal
    description: str
    subtype: str | None = None
    center_name: str = ""


@dataclass
class ReconcileRow:
    opening_id: int
    day: date
    center_name: str
    fact: Decimal
    computed: Decimal
    diff: Decimal
    is_start: bool = False


@dataclass
class CashReport:
    date_from: date
    date_to: date
    center_id: int | None
    opening: Decimal = ZERO
    cash_in: Decimal = ZERO
    card_in: Decimal = ZERO
    transfer_in: Decimal = ZERO
    promo_in: Decimal = ZERO
    expense: Decimal = ZERO
    collection: Decimal = ZERO
    by_method: dict[str, Decimal] = field(default_factory=dict)
    income_rows: list[IncomeRow] = field(default_factory=list)
    expense_rows: list[OutflowRow] = field(default_factory=list)
    collection_rows: list[OutflowRow] = field(default_factory=list)
    reconcile_rows: list[ReconcileRow] = field(default_factory=list)

    @property
    def outflow(self) -> Decimal:
        return money(self.expense + self.collection)

    @property
    def reconcile(self) -> Decimal:
        return money(sum((r.diff for r in self.reconcile_rows), ZERO))

    @property
    def has_start(self) -> bool:
        return any(r.is_start for r in self.reconcile_rows)

    @property
    def closing(self) -> Decimal:
        return money(
            self.opening + self.cash_in - self.expense - self.collection + self.reconcile
        )

    @property
    def val_rows(self) -> list[tuple[str, Decimal]]:
        return [
            (method.value, money(self.by_method.get(method.value, ZERO)))
            for method in VAL_METHODS
        ]

    @property
    def val(self) -> Decimal:
        return money(sum((amount for _, amount in self.val_rows), ZERO))


def _center_ids(db: Session, center_id: int | None) -> list[int]:
    if center_id:
        return [center_id]
    return list(db.scalars(select(ServiceCenter.id).order_by(ServiceCenter.id)))


def cash_date(order: ServiceOrder) -> date:
    """Till day of an order: pay date in Tbilisi; legacy orders without paid_at use order_date."""
    paid = as_business_date(order.paid_at)
    return paid if paid is not None else order.order_date


def _till_moves(db: Session, center_id: int, until: date) -> list[tuple[date, Decimal]]:
    """Signed cash movements of one center up to `until` inclusive."""
    moves: list[tuple[date, Decimal]] = []
    rows = db.execute(
        select(
            ServiceOrder.id,
            ServiceOrder.paid_at,
            ServiceOrder.order_date,
            func.sum(OrderLine.amount_work + OrderLine.amount_parts),
        )
        .join(OrderLine, OrderLine.order_id == ServiceOrder.id)
        .where(
            ServiceOrder.status == OrderStatus.issued,
            ServiceOrder.service_center_id == center_id,
            OrderLine.payment_type == PaymentType.cash,
        )
        .group_by(ServiceOrder.id, ServiceOrder.paid_at, ServiceOrder.order_date)
    )
    for _, paid_at, order_date, amount in rows:
        paid = as_business_date(paid_at)
        day = paid if paid is not None else order_date
        if day <= until:
            moves.append((day, money(amount)))
    entries = db.execute(
        select(CashEntry.entry_date, CashEntry.amount).where(
            CashEntry.service_center_id == center_id,
            CashEntry.entry_date <= until,
        )
    )
    for day, amount in entries:
        moves.append((day, -money(amount)))
    return moves


def _sum_moves(moves: list[tuple[date, Decimal]], start: date | None, end_excl: date) -> Decimal:
    return money(
        sum(
            (amt for day, amt in moves if (start is None or day >= start) and day < end_excl),
            ZERO,
        )
    )


def _morning_balance(
    snapshots: list[CashOpening],
    moves: list[tuple[date, Decimal]],
    day: date,
) -> Decimal:
    """Cash on the morning of `day`: last snapshot on or before it + movements since.

    Before the first "start of accounting" snapshot the till is not tracked (0).
    """
    base = None
    for snap in snapshots:
        if snap.as_of_date <= day:
            base = snap
    if base is not None:
        return money(money(base.amount) + _sum_moves(moves, base.as_of_date, day))
    if any(s.is_start for s in snapshots):
        return ZERO
    return _sum_moves(moves, None, day)


def _snapshots(db: Session, center_id: int) -> list[CashOpening]:
    return list(
        db.scalars(
            select(CashOpening)
            .where(CashOpening.service_center_id == center_id)
            .order_by(CashOpening.as_of_date)
        )
    )


def opening_balance(db: Session, center_ids: list[int], as_of: date) -> Decimal:
    total = ZERO
    for cid in center_ids:
        moves = _till_moves(db, cid, as_of)
        total += _morning_balance(_snapshots(db, cid), moves, as_of)
    return money(total)


def lock_date(db: Session, center_id: int | None) -> date | None:
    """Days before the latest reconciliation are closed."""
    if not center_id:
        return None
    return db.scalar(
        select(func.max(CashOpening.as_of_date)).where(CashOpening.service_center_id == center_id)
    )


def _center_till(
    db: Session, center_id: int, date_from: date, date_to: date
) -> tuple[Decimal, list[ReconcileRow]]:
    """Opening of the period and reconciliation rows inside it, as one running chain."""
    snaps = _snapshots(db, center_id)
    moves = _till_moves(db, center_id, date_to)
    opening = _morning_balance(snaps, moves, date_from)
    inside = [s for s in snaps if date_from < s.as_of_date <= date_to]
    rows: list[ReconcileRow] = []
    if not inside:
        return opening, rows
    center = db.get(ServiceCenter, center_id)
    bal, point = opening, date_from
    for snap in inside:
        computed = money(bal + _sum_moves(moves, point, snap.as_of_date))
        fact = money(snap.amount)
        rows.append(
            ReconcileRow(
                opening_id=snap.id,
                day=snap.as_of_date,
                center_name=center.name if center else "",
                fact=fact,
                computed=computed,
                diff=money(fact - computed),
                is_start=snap.is_start,
            )
        )
        bal, point = fact, snap.as_of_date
    return opening, rows


def build_cash_report(
    db: Session,
    date_from: date,
    date_to: date,
    center_id: int | None = None,
) -> CashReport:
    ids = _center_ids(db, center_id)
    report = CashReport(
        date_from=date_from,
        date_to=date_to,
        center_id=center_id,
    )
    if not ids:
        return report
    opening = ZERO
    for cid in ids:
        center_opening, rows = _center_till(db, cid, date_from, date_to)
        opening += center_opening
        report.reconcile_rows.extend(rows)
    report.opening = money(opening)
    report.reconcile_rows.sort(key=lambda r: (r.day, r.center_name))

    orders = list(
        db.scalars(
            select(ServiceOrder)
            .options(
                joinedload(ServiceOrder.lines),
                joinedload(ServiceOrder.assignee),
                joinedload(ServiceOrder.machine),
                joinedload(ServiceOrder.service_center),
            )
            .where(
                ServiceOrder.status == OrderStatus.issued,
                ServiceOrder.service_center_id.in_(ids),
            )
            .order_by(ServiceOrder.order_date, ServiceOrder.id)
        ).unique()
    )
    for order in orders:
        day = cash_date(order)
        if day < date_from or day > date_to:
            continue
        buckets = {method.value: ZERO for method in VAL_METHODS}
        for line in order.lines:
            if line.payment_type not in VAL_METHODS:
                continue
            buckets[line.payment_type.value] += money(line.amount_work) + money(line.amount_parts)
        if all(amount == ZERO for amount in buckets.values()):
            continue
        assignee = (
            (order.assignee.full_name or order.assignee.login) if order.assignee else ""
        )
        note = (order.machine.serial_number if order.machine else "") or f"#{order.id}"
        for method in VAL_METHODS:
            amount = money(buckets[method.value])
            if amount == ZERO:
                continue
            report.by_method[method.value] = money(
                report.by_method.get(method.value, ZERO) + amount
            )
            report.income_rows.append(
                IncomeRow(
                    order_id=order.id,
                    order_date=day,
                    payment_type=method,
                    amount=amount,
                    assignee=assignee,
                    note=note,
                )
            )
    report.cash_in = money(report.by_method.get(PaymentType.cash.value, ZERO))
    report.card_in = money(report.by_method.get(PaymentType.card.value, ZERO))
    report.transfer_in = money(report.by_method.get(PaymentType.transfer.value, ZERO))
    report.promo_in = money(report.by_method.get(PaymentType.promo.value, ZERO))
    for method in VAL_METHODS:
        report.by_method.setdefault(method.value, ZERO)

    entries = list(
        db.scalars(
            select(CashEntry)
            .options(joinedload(CashEntry.service_center))
            .where(
                CashEntry.entry_date >= date_from,
                CashEntry.entry_date <= date_to,
                CashEntry.service_center_id.in_(ids),
            )
            .order_by(CashEntry.entry_date, CashEntry.id)
        )
    )
    for entry in entries:
        row = OutflowRow(
            entry_id=entry.id,
            entry_date=entry.entry_date,
            amount=money(entry.amount),
            description=entry.description,
            subtype=entry.subtype.value if entry.subtype else None,
            center_name=entry.service_center.name if entry.service_center else "",
        )
        if entry.kind == CashEntryKind.expense:
            report.expense += row.amount
            report.expense_rows.append(row)
        elif entry.kind == CashEntryKind.collection:
            report.collection += row.amount
            report.collection_rows.append(row)
    report.expense = money(report.expense)
    report.collection = money(report.collection)
    return report


def recent_cash_entries(
    db: Session, center_id: int | None = None, limit: int = 30
) -> list[CashEntry]:
    stmt = (
        select(CashEntry)
        .options(joinedload(CashEntry.service_center))
        .order_by(CashEntry.entry_date.desc(), CashEntry.id.desc())
        .limit(limit)
    )
    if center_id:
        stmt = stmt.where(CashEntry.service_center_id == center_id)
    return list(db.scalars(stmt).unique())


def cash_report_xlsx(report: CashReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cash"
    period = f"{report.date_from.isoformat()} — {report.date_to.isoformat()}"
    ws["A1"] = f"Period {period}"
    ws["A2"] = "Opening"
    ws["B2"] = float(report.opening)
    ws["A3"] = "Cash in"
    ws["B3"] = float(report.cash_in)
    ws["A4"] = "Terminal"
    ws["B4"] = float(report.card_in)
    ws["A5"] = "Expense"
    ws["B5"] = float(report.expense)
    ws["A6"] = "Collection"
    ws["B6"] = float(report.collection)
    ws["A7"] = "Closing"
    ws["B7"] = float(report.closing)
    ws["A8"] = (
        "opening + cash - expense - collection + reconcile (terminal, transfer and promo not in till)"
    )
    ws["A9"] = "Transfer"
    ws["B9"] = float(report.transfer_in)
    excel_labels = {
        PaymentType.cash.value: "Cash",
        PaymentType.card.value: "Terminal",
        PaymentType.transfer.value: "Transfer",
        PaymentType.promo.value: "Promo",
    }
    row_no = 11
    ws.cell(row_no, 1, "Val")
    row_no += 1
    for method_value, amount in report.val_rows:
        ws.cell(row_no, 1, excel_labels.get(method_value, method_value))
        ws.cell(row_no, 2, float(amount))
        row_no += 1
    ws.cell(row_no, 1, "Val total")
    ws.cell(row_no, 2, float(report.val))
    if report.reconcile_rows:
        row_no += 2
        ws.cell(row_no, 1, "Reconcile")
        ws.cell(row_no, 2, float(report.reconcile))
        row_no += 1
        for col, title in enumerate(
            ("Date", "Center", "Fact", "Computed", "Difference", "Start"), 1
        ):
            ws.cell(row_no, col, title)
        for r in report.reconcile_rows:
            row_no += 1
            ws.cell(row_no, 1, r.day.isoformat())
            ws.cell(row_no, 2, r.center_name)
            ws.cell(row_no, 3, float(r.fact))
            ws.cell(row_no, 4, float(r.computed))
            ws.cell(row_no, 5, float(r.diff))
            ws.cell(row_no, 6, "yes" if r.is_start else "")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
