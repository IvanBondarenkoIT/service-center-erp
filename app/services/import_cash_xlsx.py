from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    CashEntry,
    CashEntryKind,
    CashOpening,
    CollectionSubtype,
    ServiceCenter,
)
from app.services.cash_book import cash_entry_key, existing_cash_entry, money

_DATE_SPLIT = re.compile(r"[,./\-\s]+")


@dataclass
class _SheetLayout:
    expense_date_col: int = 8
    expense_amount_col: int = 9
    expense_desc_col: int = 10
    collection_date_col: int = 12
    collection_amount_col: int = 13
    collection_whom_col: int = 14
    income_cash_col: int = 2
    income_terminal_col: int = 4
    day_col: int = 1
    opening_amount_col: int | None = 18


@dataclass
class CashImportStats:
    sheets: int = 0
    expenses_created: int = 0
    expenses_skipped: int = 0
    collections_created: int = 0
    collections_skipped: int = 0
    openings_upserted: int = 0
    income_rows_ignored: int = 0

    def as_dict(self) -> dict[str, int]:
        return {
            "sheets": self.sheets,
            "expenses_created": self.expenses_created,
            "expenses_skipped": self.expenses_skipped,
            "collections_created": self.collections_created,
            "collections_skipped": self.collections_skipped,
            "openings_upserted": self.openings_upserted,
            "income_rows_ignored": self.income_rows_ignored,
        }


def _norm(value: Any) -> str:
    return str(value or "").strip().lower()


def _to_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text or text in {".", "-"}:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def parse_cash_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    parts = [p for p in _DATE_SPLIT.split(text) if p]
    if len(parts) != 3:
        return None
    try:
        day, month, year = (int(parts[0]), int(parts[1]), int(parts[2]))
    except ValueError:
        return None
    if year < 100:
        year += 2000
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _map_collection_subtype(raw: Any) -> CollectionSubtype:
    text = _norm(raw)
    if any(token in text for token in ("аванс", "advance", "ავანს")):
        return CollectionSubtype.advance
    if any(
        token in text
        for token in ("зарплат", "salary", "xelfas", "ხელფას")
    ):
        return CollectionSubtype.salary
    return CollectionSubtype.bank


def _cell_label(value: Any) -> str:
    return _norm(value).replace("ё", "е")


def _detect_layout(ws: Worksheet) -> _SheetLayout:
    layout = _SheetLayout()
    row1 = {_cell_label(ws.cell(1, col).value): col for col in range(1, 21)}
    row2 = {_cell_label(ws.cell(2, col).value): col for col in range(1, 21)}

    expense_header = row1.get("расход")
    if expense_header:
        layout.expense_amount_col = expense_header
        layout.expense_date_col = expense_header - 1
        layout.expense_desc_col = expense_header + 1

    collection_header = row1.get("инкасация")
    if collection_header:
        layout.collection_amount_col = collection_header
        layout.collection_date_col = collection_header - 1
        layout.collection_whom_col = collection_header + 1

    for label, col in row2.items():
        if "касса /cash" in label or label in {"касса /cash", "касса / cash"}:
            layout.income_cash_col = col
        if "терминал" in label:
            layout.income_terminal_col = col
        if label == "на что":
            layout.expense_desc_col = col
        if label == "кому":
            layout.collection_whom_col = col

    layout.opening_amount_col = None
    for col in range(1, 21):
        label = _cell_label(ws.cell(1, col).value)
        if label != "касса":
            continue
        for neighbor in (col + 1, col - 1):
            if neighbor < 1:
                continue
            if _to_decimal(ws.cell(1, neighbor).value) is not None:
                layout.opening_amount_col = neighbor
                break
        if layout.opening_amount_col:
            break
    return layout


def _find_center(db: Session, center_code: str) -> ServiceCenter:
    center = db.scalar(select(ServiceCenter).where(ServiceCenter.code == center_code))
    if center is None:
        raise ValueError(f"Unknown service center: {center_code}")
    return center


def _upsert_opening(
    db: Session,
    *,
    center_id: int,
    as_of: date,
    amount: Decimal,
) -> None:
    existing = db.scalar(
        select(CashOpening).where(
            CashOpening.service_center_id == center_id,
            CashOpening.as_of_date == as_of,
        )
    )
    if existing:
        existing.amount = amount
        return
    db.add(
        CashOpening(
            service_center_id=center_id,
            as_of_date=as_of,
            amount=amount,
        )
    )


def _opening_as_of(ws: Worksheet, layout: _SheetLayout) -> date | None:
    for row in range(3, (ws.max_row or 3) + 1):
        parsed = parse_cash_date(ws.cell(row, layout.day_col).value)
        if parsed:
            return parsed
    return None


def _iter_data_rows(ws: Worksheet):
    max_row = ws.max_row or 0
    for row in range(3, max_row + 1):
        yield row


def import_cash_xlsx(
    db: Session,
    path: Path,
    *,
    center_code: str = "batumi",
    created_by_id: int | None = None,
) -> dict[str, int]:
    center = _find_center(db, center_code)
    wb = load_workbook(path, data_only=True)
    stats = CashImportStats()
    seen: set[tuple] = set()

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        layout = _detect_layout(ws)
        stats.sheets += 1
        sheet_entries = 0

        for row in _iter_data_rows(ws):
            day = parse_cash_date(ws.cell(row, layout.day_col).value)
            cash_in = _to_decimal(ws.cell(row, layout.income_cash_col).value)
            terminal_in = _to_decimal(ws.cell(row, layout.income_terminal_col).value)
            if (cash_in and cash_in != 0) or (terminal_in and terminal_in != 0):
                stats.income_rows_ignored += 1

            exp_amount = _to_decimal(ws.cell(row, layout.expense_amount_col).value)
            exp_desc = str(ws.cell(row, layout.expense_desc_col).value or "").strip()[:512]
            if exp_amount and exp_amount > 0:
                exp_date = parse_cash_date(ws.cell(row, layout.expense_date_col).value) or day
                if exp_date is None:
                    stats.expenses_skipped += 1
                else:
                    amount = money(exp_amount)
                    key = cash_entry_key(
                        center.id, exp_date, CashEntryKind.expense, amount, exp_desc
                    )
                    if key in seen or existing_cash_entry(
                        db,
                        center_id=center.id,
                        entry_date=exp_date,
                        kind=CashEntryKind.expense,
                        amount=amount,
                        description=exp_desc,
                    ):
                        stats.expenses_skipped += 1
                    else:
                        db.add(
                            CashEntry(
                                entry_date=exp_date,
                                service_center_id=center.id,
                                kind=CashEntryKind.expense,
                                subtype=None,
                                amount=amount,
                                description=exp_desc,
                                created_by_id=created_by_id,
                            )
                        )
                        seen.add(key)
                        stats.expenses_created += 1
                        sheet_entries += 1

            col_amount = _to_decimal(ws.cell(row, layout.collection_amount_col).value)
            whom_raw = ws.cell(row, layout.collection_whom_col).value
            whom = str(whom_raw or "").strip()[:512]
            if col_amount and col_amount > 0:
                col_date = parse_cash_date(ws.cell(row, layout.collection_date_col).value) or day
                if col_date is None:
                    stats.collections_skipped += 1
                else:
                    amount = money(col_amount)
                    key = cash_entry_key(
                        center.id, col_date, CashEntryKind.collection, amount, whom
                    )
                    if key in seen or existing_cash_entry(
                        db,
                        center_id=center.id,
                        entry_date=col_date,
                        kind=CashEntryKind.collection,
                        amount=amount,
                        description=whom,
                    ):
                        stats.collections_skipped += 1
                    else:
                        db.add(
                            CashEntry(
                                entry_date=col_date,
                                service_center_id=center.id,
                                kind=CashEntryKind.collection,
                                subtype=_map_collection_subtype(whom_raw),
                                amount=amount,
                                description=whom,
                                created_by_id=created_by_id,
                            )
                        )
                        seen.add(key)
                        stats.collections_created += 1
                        sheet_entries += 1

        opening_raw = (
            ws.cell(1, layout.opening_amount_col).value
            if layout.opening_amount_col
            else None
        )
        opening_amount = _to_decimal(opening_raw)
        as_of = _opening_as_of(ws, layout)
        if opening_amount is not None and as_of is not None:
            quantized = money(opening_amount)
            if quantized != 0 or sheet_entries:
                _upsert_opening(
                    db,
                    center_id=center.id,
                    as_of=as_of,
                    amount=quantized,
                )
                stats.openings_upserted += 1

    db.commit()
    return stats.as_dict()
